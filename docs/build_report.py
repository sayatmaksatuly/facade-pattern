"""Build the formal report from the completed project and saved verification evidence.

Requires python-docx. The final report is rendered and inspected separately.
"""
from pathlib import Path
import json
import re
import subprocess
from html import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = json.loads((ROOT / 'docs/evidence/verification.json').read_text(encoding='utf-8-sig'))
TRANSCRIPT = (ROOT / 'docs/evidence/cli-verification.txt').read_text(encoding='utf-8-sig')
QA = ROOT / 'docs/.qa'
QA.mkdir(parents=True, exist_ok=True)

def compact_diagram():
    elements = ['<svg xmlns="http://www.w3.org/2000/svg" width="1000" height="655" viewBox="0 0 1000 655">',
        '<defs><marker id="a" markerWidth="7" markerHeight="7" refX="6" refY="3.5" orient="auto"><path d="M0 0 L6 3.5 L0 7" fill="none" stroke="#454545"/></marker></defs>',
        '<rect width="1000" height="655" fill="white"/>',
        '<style>text{font-family:Arial,sans-serif;fill:#111}.n{font-size:19px;font-weight:bold}.l{font-size:16px}</style>']
    def arrow(points, dashed=False):
        dash = ' stroke-dasharray="6 4"' if dashed else ''
        elements.append(f'<polyline points="{points}" fill="none" stroke="#454545" stroke-width="1.5" marker-end="url(#a)"{dash}/>')
    def box(x,y,w,h,label,stereo=None):
        elements.extend([f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="white" stroke="#454545" stroke-width="1.4"/>',
                         f'<text x="{x+w/2}" y="{y+28}" text-anchor="middle" class="n">{escape(label)}</text>'])
    groups = [('A gate', ['GateController','LicensePlateReader','GateGateway','GateAdmission']),
              ('B spaces', ['ReservationManager','SpaceSensorGrid','ReservationGateway','ParkingReservation']),
              ('C payment', ['PaymentTerminal','OverstayAlert','PaymentGateway','ParkingPayment'])]
    for i,(label,members) in enumerate(groups):
        x = 20+330*i
        elements += [f'<rect x="{x}" y="200" width="300" height="330" rx="4" fill="#f7f7f7" stroke="#aaa"/>',
                     f'<text x="{x+150}" y="222" text-anchor="middle" class="n">{label}</text>']
    arrow('450,70 450,85 300,85 300,105')
    arrow('550,70 550,85 700,85 700,105')
    arrow('550,130 450,130')
    for sx,rail,offset in [(300,170,-110),(700,185,110)]:
        for target in [170,500,830]:
            arrow(f'{sx},150 {sx},{rail} {target+offset},{rail} {target+offset},235')
    for i in range(3):
        x=35+330*i
        center=x+135
        arrow(f'{center},280 {center},310')
        arrow(f'{x+10},280 {x-8},280 {x-8},405 {x},405')
        arrow(f'{center},430 {center},460', True)
        arrow(f'{x+270},405 {x+280},405 {x+280},555 775,555 775,585')
    arrow('635,385 660,385 660,333 635,333', True)
    arrow('150,128 8,128 8,606 65,606')
    arrow('850,128 992,128 992,545 48,545 48,615 65,615')
    arrow('355,607 645,607')
    box(365,25,270,45,'ParkingConsole')
    box(150,105,300,45,'ParkingFacade')
    box(550,105,300,45,'AdministrationFacade')
    for i,(_,members) in enumerate(groups):
        for y,name in zip([235,310,385,460],members): box(35+330*i,y,270,45,name)
    box(65,585,290,45,'WorkflowStore')
    box(645,585,290,45,'Database')
    elements.append('<text x="500" y="650" text-anchor="middle" class="l">Solid arrows show collaborators; dashed arrows show use dependencies. Gateways persist their own entities.</text></svg>')
    source=QA/'uml-report.svg'
    source.write_text('\n'.join(elements),encoding='utf-8')
    deps=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies'
    node=deps/'node/bin/node.exe'
    sharp=deps/'node/node_modules/sharp'
    js="require("+json.dumps(str(sharp))+ ")("+json.dumps(str(source))+").png().toFile("+json.dumps(str(QA/'uml-report.png'))+");"
    subprocess.run([str(node),'-e',js],check=True)

compact_diagram()
doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.27);section.page_height=Inches(11.69)
section.top_margin=Inches(.67);section.bottom_margin=Inches(.67)
section.left_margin=Inches(.72);section.right_margin=Inches(.72)
section.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    style=doc.styles[name]
    style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(14)
doc.styles['Normal'].paragraph_format.line_spacing=1.05
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(26)
doc.styles['Title'].paragraph_format.space_after=Pt(11)
doc.styles['Subtitle'].font.size=Pt(15)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
doc.styles['Heading 2'].font.size=Pt(14)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.keep_with_next=True
footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text, style=None):
    return doc.add_paragraph(text,style)
def h(text):return doc.add_heading(text,level=1)
def page():doc.add_page_break()
def code(text):
    paragraph=doc.add_paragraph()
    paragraph.paragraph_format.space_before=Pt(3)
    paragraph.paragraph_format.space_after=Pt(8)
    paragraph.paragraph_format.line_spacing=1
    paragraph.paragraph_format.keep_together=True
    run=paragraph.add_run(text)
    run.font.name='Consolas';run.font.size=Pt(10)
    return paragraph
def table(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for cell,width in zip(t.rows[0].cells,widths):cell.width=Inches(width)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=text
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):cells[i].text=str(text);cells[i].width=Inches(widths[i])
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            tc=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6' if ri==0 else 'FFFFFF');tc.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            tc.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(2)
                paragraph.paragraph_format.line_spacing=1
                for run in paragraph.runs:run.font.size=Pt(11);run.bold=ri==0
    p('')
    return t

doc.add_paragraph('Smart Parking Garage',style='Title')
doc.add_paragraph('Software Design Patterns Assignment 4',style='Subtitle')
p('Facade over a three stage chain with durable persistence and recovery')
p('Project: smart-parking 1.0.0\nAssigned variant: 20\nVerification date: 16 October 2026\nStack: Java 17, Maven, JUnit 5, JaCoCo and SQLite')
h('1 Introduction and objective')
p('The project confirms a prepaid parking booking through three independent subsystems. A single client request holds an entry clearance, reserves a space and records payment. If a later stage fails after earlier commits, the facade reads persisted state and compensates the effects that actually exist.')
p(f'The completed implementation passed {EVIDENCE["tests"]} tests with no failures, errors or skipped tests. The required coverage scope reached {EVIDENCE["required_scope"]["percent"]:.2f}% line coverage. A separate process demonstration verified both compensation depths and recovery after forced termination. These results support the implementation; they do not predict an assessment grade.')
h('2 Assigned domain and assessment scope')
p('The source assignment identifies Smart Parking Garage as variant 20: Gate Controller and License Plate Reader in A; Space Sensor Grid and Reservation Manager in B; Payment Terminal and Overstay Alert in C. Those boundaries and starting points are retained.')
p('The use case is a prepaid booking for one of three spaces. The rate is 500 cents per hour for 1–24 hours, with a maximum vehicle height of 220 cm. A completed booking retains a HELD entry clearance and a RESERVED space for later physical admission. Payment is a durable local ledger entry; no bank or hardware is contacted.')
table(['Assessment component','Weight'],[['Defense and live coding','50%'],['Unit tests','25%'],['Working demonstration','10%'],['Documentation','10%'],['Subsystems facades and persistence','5%']],[5.5,1.1])

page();h('3 Architecture and SOLID')
p('ParkingConsole is the client. It handles input and output through two facades without importing a subsystem or JDBC API. ParkingApplication manages process startup and a database file lock. GarageRuntime is the composition root that constructs the independent components and injects their dependencies through constructors.')
p('Single Responsibility is visible in the separation of entity transitions, pure business rules, persistence gateways, workflow coordination and console interaction. The Principle of Least Knowledge limits the client to its immediate facade collaborators. The implementation uses small concrete classes where only one local implementation exists. Constructor injection exposes dependencies; it does not by itself prove Dependency Inversion. A future hardware boundary would justify a dedicated interface.')
h('4 Three independent subsystems')
p('A contains GateAdmission, LicensePlateReader, GateGateway and GateController. LicensePlateReader normalizes and validates a plate and checks vehicle clearance. GateAdmission owns HELD, ADMITTED and CANCELLED, allowing admission only from HELD and preserving a cancelled record. GateGateway directly persists the permit.')
p('B contains ParkingReservation, SpaceSensorGrid, ReservationGateway and ReservationManager. SpaceSensorGrid decides whether a numbered space is usable from persisted obstruction and allocation data. ParkingReservation owns RESERVED, OCCUPIED and RELEASED. The gateway checks and stores its own allocation without calling A or C.')
p('C contains ParkingPayment, OverstayAlert, PaymentGateway and PaymentTerminal. OverstayAlert calculates the charge, validates tender and computes overdue minutes. ParkingPayment owns AUTHORIZED, CAPTURED and REFUNDED. PaymentGateway stores the local payment effect and supports capture and refund.')
p('Each subsystem has four top-level classes, for twelve in total. Nested status enums are part of their entity. Every subsystem has a state entity, independent pure logic and its own SQL gateway. None imports another subsystem or a facade. Entity methods reject illegal transitions independently; gateways make repeated recovery requests safe by recognizing already committed outcomes.')
p('The standalone command and menu exercise the public gate, spaces or payment API in isolation. The corresponding JUnit tests construct only the selected subsystem and a fresh SQLite store.')

page();h('5 Main and administrative facades')
p('ParkingFacade.park(ParkingRequest) is the ordinary client operation. It persists or reloads the request and calls resume. The resume loop reloads workflow progress before choosing the next stage. advance sequences gate.hold, spaces.reserve and payments.charge. handleFailure inspects persisted entities before deciding whether to reject or compensate.')
p('AdministrationFacade exposes inspection, event history, sensor maintenance, recovery and manual cancellation. recoverAll processes incomplete operations separately so one unavailable operation does not prevent the others from being considered. The common facade remains usable without the CLI.')
code('Workflow result = parking.park(\n    new ParkingRequest("booking-1", "ABC123", 180, 1, 2, 1000));')
h('6 Durable SQLite persistence')
p('Each subsystem owns its tables through its gateway. WorkflowStore owns request input, progress and the ordered event journal. Database provides a JDBC connection and a small transaction helper. Foreign keys, WAL journaling, a busy timeout and synchronous FULL are enabled.')
p('Each gateway action commits independently. The complete A to B to C workflow is deliberately not one database transaction. A workflow progress transition and its associated event commit together, but a domain commit and the following progress update are separate. That interruption window is handled by idempotent operations identified by the original operation ID.')
table(['Persisted state','Meaning'],[['PENDING_A','Request exists and A must be reconciled'],['PENDING_B','A recorded and B is next'],['PENDING_C','A and B recorded and C is next'],['COMPLETED','Prepaid booking confirmed'],['COMPENSATING','Durable intent to reverse effects'],['COMPENSATED','Required reversals completed'],['REJECTED','Failure before any domain commit']],[1.7,4.9])
p('A repeated identical request returns its existing result; a different request under the same ID is rejected. Compensation retains records with terminal statuses. If compensation itself fails, COMPENSATING remains available for a later administrative retry.')

page();h('7 Successful workflow')
p('For a valid plate, a vehicle within clearance, a clear free space and enough tender, A commits HELD, B commits RESERVED and C commits CAPTURED. The journal then records COMPLETED. The verified example used two hours and 1000 cents.')
code('3. A committed: entry clearance HELD; next B.\n5. B committed: space RESERVED; next C.\n7. C committed: payment CAPTURED; workflow COMPLETED.\nFinal persisted state: COMPLETED')
h('8 Failure at B and compensation of A')
p('The demonstration persists a blocked sensor reading for space 2. A commits before B reads that condition. B rejects the reservation through its normal availability rule. The facade persists compensation intent, cancels A and finishes COMPENSATED. B and C have no record for the rejected operation.')
code('10. A committed: entry clearance HELD; next B.\n11. Stage B executing: check sensor and reserve space.\n13. Compensation A committed: entry clearance CANCELLED.\nFinal persisted state: COMPENSATED')
h('9 Failure at C and compensation of B then A')
p('The C demonstration supplies 100 cents for a two-hour booking costing 1000 cents. A and B have already committed when the ordinary payment rule rejects the tender. The persisted history below shows B released before A is cancelled. No payment record is created.')
code('17. A committed: entry clearance HELD; next B.\n19. B committed: space RESERVED; next C.\n22. Compensation B committed: space RELEASED.\n23. Compensation A committed: entry clearance CANCELLED.\nFinal persisted state: COMPENSATED')
p('These are selected lines from docs/evidence/cli-verification.txt. The full transcript includes failure reasons and later database inspection. A separate oversized-vehicle scenario fails at A before domain effects and is persisted as REJECTED.')

page();h('10 Restart and recovery demonstration')
p('The process harness launched the actual packaged application, entered the interactive checkpoint command and waited until A committed and PENDING_B was printed. A second process was denied access by the OS file lock. The harness forcibly killed the original JVM, then opened the same SQLite file in a new JVM. Inspection found the HELD permit and PENDING_B progress before recovery.')
code('Final persisted state: PENDING_B\nThis database is already open in another garage process.\nExit code: 2\nOriginal Java process was forcibly killed after A committed. OS lock released.')
p('The new process ran recover crash-booking, committed B and C, and produced COMPLETED. The final inspection retained the same operation ID and showed HELD, RESERVED and CAPTURED. This demonstrates a real termination and disk reload rather than reuse of in-memory objects.')
p('Recovery also covers crashes between a subsystem commit and its journal marker. The facade repeats the stage with the original ID; the gateway recognizes the existing matching record. Compensation examines persisted C, B and A in reverse order. If undoing B fails, A is not cancelled prematurely. Recovery later retries from the recorded states. Store unavailability still requires restoring access before any recovery can proceed.')
h('11 F I R S T unit testing')
p('Fast tests use local SQLite files without network calls. Independent tests receive fresh temporary directories and known persisted state. Repeatable tests use fixed inputs and do not rely on execution order. Self-validating tests assert statuses, record presence and event order. Production code and its tests are retained in genuine development commits; no earlier history has been fabricated.')
p('The domain-specific helpers aDriverRequestsParking, entryAndReservationHaveCommitted, theGarageRestarts and assertBookingIsDurablyConfirmed make the workflow cases read in parking terms. Mandatory B and C failure tests close and reopen the database before asserting the final state. Additional cases cover commit-to-journal gaps, interrupted compensation, repeated cancellation, conflicting operation IDs, invalid transitions and standalone subsystem behavior.')

page();h('12 Actual tests and JaCoCo coverage')
p(f'The command .\\mvnw.cmd clean verify completed successfully on 16 October 2026. Surefire reported {EVIDENCE["tests"]} tests across {len(EVIDENCE["test_suites"])} suites, {EVIDENCE["failures"]} failures, {EVIDENCE["errors"]} errors and {EVIDENCE["skipped"]} skipped tests. The values below are read from the saved verification summary generated from Surefire and JaCoCo XML.')
rows=[]
for name,values in EVIDENCE['required_coverage_targets'].items():
    rows.append([name,f'{values["covered"]} / {values["total"]}',f'{values["percent"]:.2f}%'])
for label,key in [('Combined required scope','required_scope'),('Whole project','whole_project')]:
    v=EVIDENCE[key];rows.append([label,f'{v["covered"]} / {v["total"]}',f'{v["percent"]:.2f}%'])
table(['Coverage target','Covered lines','Line coverage'],rows,[3.4,1.6,1.6])
p('The required-scope percentage is weighted by executable lines across the three subsystem packages and the two facade classes. Maven enforces at least 80% separately for each of those five targets. An intentionally impossible 101% threshold rejected all five, confirming that the configured checks select them.')
p('The HTML report is target/site/jacoco/index.html; detailed counters are in target/site/jacoco/jacoco.xml. Test reports are in target/surefire-reports. The saved summary is docs/evidence/verification.json, reproduced by scripts/export-verification.ps1.')
suite_rows=[]
suites=EVIDENCE['test_suites']
for a,b in zip(suites[:4],suites[4:]):
    suite_rows.append([a['name'].split('.')[-1],a['tests'],b['name'].split('.')[-1],b['tests']])
table(['Test suite','Tests','Test suite','Tests'],suite_rows,[2.75,.55,2.75,.55])
p('Coverage shows executed lines, not proof that every branch and business scenario is correct. The persisted-state assertions and ordered compensation assertions are the behavioral evidence. The CLI harness separately validates the packaged application and its process lifecycle.')

land=doc.add_section(WD_SECTION_START.NEW_PAGE)
land.orientation=WD_ORIENT.LANDSCAPE;land.page_width=Inches(11.69);land.page_height=Inches(8.27)
land.top_margin=Inches(.5);land.bottom_margin=Inches(.5);land.left_margin=Inches(.7);land.right_margin=Inches(.7)
h('13 UML class diagram')
caption=p('Client, facades, all twelve subsystem classes and persistence dependencies. The detailed UML with methods, states and composition root is in docs/uml.svg; editable source is docs/uml.puml.')
for run in caption.runs:run.font.size=Pt(12)
paragraph=doc.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
paragraph.paragraph_format.space_after=Pt(0)
paragraph.add_run().add_picture(str(QA/'uml-report.png'),width=Inches(9.2))
portrait=doc.add_section(WD_SECTION_START.NEW_PAGE)
portrait.orientation=WD_ORIENT.PORTRAIT;portrait.page_width=Inches(8.27);portrait.page_height=Inches(11.69)
portrait.top_margin=Inches(.67);portrait.bottom_margin=Inches(.67);portrait.left_margin=Inches(.72);portrait.right_margin=Inches(.72)
h('14 Principle of Least Knowledge')
p('Without a facade, the booking client must know three subsystem APIs and their order. It would also need to record progress, detect partial success and organize reversals. This illustrative before fragment uses the actual public APIs:')
code('void bookDirectly(GateController gate, ReservationManager spaces,\n                  PaymentTerminal payments, ParkingRequest r) {\n    gate.hold(r.operationId(), r.plate(), r.heightCm());\n    spaces.reserve(r.operationId(), r.space());\n    payments.charge(r.operationId(), r.hours(), r.tenderCents());\n    // Client would also own progress and compensation.\n}')
p('The implemented client has one immediate workflow collaborator. It supplies immutable input and receives a workflow result:')
code('Workflow book(ParkingFacade parking, ParkingRequest request) {\n    return parking.park(request);\n}')
p('The facade owns sequencing and recovery. Administrative tasks use a separate facade. The composition root is allowed to know constructors, while the console does not navigate controllers, gateways or JDBC connections. An architecture test guards the client imports and subsystem boundaries.')
h('15 Facade versus Adapter reflection')
p('A Facade presents a coherent high-level operation over existing capabilities. Here, parking.park coordinates entry clearance, a space reservation and payment, including compensation after partial success. Those independent subsystem APIs remain available on their own.')
p('An Adapter translates an interface so a component can fit an expected contract. A future adapter could translate a gate manufacturer’s protocol into the garage’s gate API. This assignment has no incompatible external interface to translate. The problem is coordinating existing operations and reducing client knowledge, so Facade matches its intent. An adapter at a future hardware boundary could coexist with the facades.')

page();h('16 Running the application and demonstration')
p('On Windows, run the one-time scripts/setup-java.ps1 if a Java 17 JDK is not already configured. The checked-in Maven Wrapper uses the project-local JDK when present. Its first build downloads dependencies; the application itself operates entirely locally.')
code('.\\mvnw.cmd -q compile exec:java\n.\\mvnw.cmd clean verify\npowershell -NoProfile -ExecutionPolicy Bypass -File scripts\\verify-demo.ps1')
p('Use menu 1 for real booking input, menu 2 for B failure and menu 3 for C failure. Menus 4 and 5 inspect persisted state and history. Menu 6 recovers incomplete operations; menu 7 performs manual cancellation; menu 8 exercises one subsystem; menu 9 leaves a durable checkpoint after A; menu 10 configures a sensor; menu 0 exits without deleting data.')
p('For the restart demonstration, choose menu 9 with a free space, terminate after PENDING_B appears, then reopen the same database and inspect it before choosing recovery. For a fresh demonstration, use a new --db path. To reset an existing demo store, stop the application and remove only that file and its -wal and -shm sidecars as documented in README.')
h('17 Defense and live coding preparation')
p('DEFENSE_GUIDE.md explains the design in Russian, identifies the main classes and methods, supplies a live demonstration script, answers twenty realistic examiner questions and proposes five live-coding exercises. The exercises identify the production files and tests for changes to clearance, space validation, payment duration, compensation failure and tariff policy.')
p('The source assignment and docs/REQUIREMENTS.md provide the requirement traceability. The student must confirm that the domain does not repeat an earlier assignment or another student’s work, publish a real repository URL and submit it through Moodle. These external submission steps have not been claimed as completed.')
h('18 Conclusion')
p('The implementation demonstrates two facades over three independently usable subsystems, durable SQLite state, both required compensation depths and recovery after a real process termination. Tests verify the persisted outcomes and line coverage exceeds each required threshold. The local model deliberately stops at prepaid booking and ledger behavior; integration with actual garage equipment, a bank and multi-process workflow execution remains outside its scope.')

# Enforce actual excerpt provenance before producing a submission document.
for expected in [
    '13. Compensation A committed: entry clearance CANCELLED.',
    '22. Compensation B committed: space RELEASED.',
    '23. Compensation A committed: entry clearance CANCELLED.',
    'Original Java process was forcibly killed after A committed. OS lock released.',
]:
    assert expected in TRANSCRIPT, f'Missing evidence for report excerpt: {expected}'
doc.core_properties.title='Smart Parking Garage Assignment 4 Report'
doc.core_properties.subject='Facade pattern with durable compensation and recovery'
doc.core_properties.author=''
doc.core_properties.keywords='Java 17, Facade, SQLite, JUnit 5, JaCoCo'
# Remove template theme overrides and title rules; keep requested academic type.
for root in [doc._element,doc.styles.element]:
    for element in list(root.iter()):
        if element.tag==qn('w:pBdr'):
            element.getparent().remove(element)
        elif element.tag==qn('w:rFonts'):
            for key in list(element.attrib):
                if 'Theme' in key:del element.attrib[key]
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run.font.name!='Consolas':run.font.name='Times New Roman'
        run.font.color.rgb=RGBColor(0,0,0)
output=ROOT/'Assignment4_Report.docx'
doc.save(output)
print(f'Created {output}')
